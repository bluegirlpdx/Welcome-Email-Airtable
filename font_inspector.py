#!/usr/bin/env python3
"""
Font Inspector Tool
Analyzes PowerPoint files to detect all fonts used and provides mapping suggestions
"""

import sys
import os
from collections import defaultdict
from pptx import Presentation
from typing import Dict, List, Set
import json


class FontInspector:
    """Inspects PowerPoint files for font usage"""

    def __init__(self, pptx_path: str):
        self.pptx_path = pptx_path
        self.presentation = None
        self.fonts_found = defaultdict(lambda: {
            'count': 0,
            'sizes': set(),
            'colors': set(),
            'bold_count': 0,
            'italic_count': 0,
            'slides': set()
        })

    def inspect(self):
        """Inspect the PowerPoint file for fonts"""
        try:
            self.presentation = Presentation(self.pptx_path)
            print(f"📄 Inspecting: {self.pptx_path}")
            print(f"📊 Total slides: {len(self.presentation.slides)}\n")

            for slide_num, slide in enumerate(self.presentation.slides, 1):
                self._inspect_slide(slide, slide_num)

            return self.get_report()

        except Exception as e:
            print(f"❌ Error inspecting file: {e}")
            return None

    def _inspect_slide(self, slide, slide_num: int):
        """Inspect a single slide for fonts"""
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.text.strip():
                            self._record_font(run, slide_num)

    def _record_font(self, run, slide_num: int):
        """Record font information from a text run"""
        font_name = run.font.name or "Unknown"

        # Increment count
        self.fonts_found[font_name]['count'] += 1
        self.fonts_found[font_name]['slides'].add(slide_num)

        # Record size
        if run.font.size:
            self.fonts_found[font_name]['sizes'].add(run.font.size.pt)

        # Record color
        try:
            if run.font.color and hasattr(run.font.color, 'rgb'):
                rgb = run.font.color.rgb
                color_hex = f"#{rgb[0]:02x}{rgb[1]:02x}{rgb[2]:02x}"
                self.fonts_found[font_name]['colors'].add(color_hex)
        except:
            pass

        # Record bold/italic
        if run.font.bold:
            self.fonts_found[font_name]['bold_count'] += 1
        if run.font.italic:
            self.fonts_found[font_name]['italic_count'] += 1

    def get_report(self) -> str:
        """Generate detailed font report"""
        lines = []
        lines.append("=" * 70)
        lines.append("FONT INSPECTION REPORT")
        lines.append("=" * 70)
        lines.append("")

        if not self.fonts_found:
            lines.append("⚠️  No fonts detected in this PowerPoint file.")
            return '\n'.join(lines)

        lines.append(f"📊 Total unique fonts found: {len(self.fonts_found)}")
        lines.append("")
        lines.append("=" * 70)
        lines.append("")

        # Sort fonts by usage count
        sorted_fonts = sorted(
            self.fonts_found.items(),
            key=lambda x: x[1]['count'],
            reverse=True
        )

        for idx, (font_name, data) in enumerate(sorted_fonts, 1):
            lines.append(f"{idx}. {font_name}")
            lines.append(f"   {'─' * 65}")
            lines.append(f"   Usage count: {data['count']} times")
            lines.append(f"   Used on slides: {', '.join(map(str, sorted(data['slides'])))}")

            if data['sizes']:
                sizes = ', '.join([f"{s:.0f}pt" for s in sorted(data['sizes'])])
                lines.append(f"   Font sizes: {sizes}")

            if data['bold_count'] or data['italic_count']:
                styles = []
                if data['bold_count']:
                    styles.append(f"Bold ({data['bold_count']}x)")
                if data['italic_count']:
                    styles.append(f"Italic ({data['italic_count']}x)")
                lines.append(f"   Styles: {', '.join(styles)}")

            if data['colors']:
                colors = ', '.join(sorted(data['colors']))
                lines.append(f"   Colors: {colors}")

            # Suggest mapping
            suggestion = self._suggest_mapping(font_name)
            if suggestion:
                lines.append(f"   💡 Suggestion: {suggestion}")

            lines.append("")

        # Summary of recommendations
        lines.append("=" * 70)
        lines.append("RECOMMENDATIONS")
        lines.append("=" * 70)
        lines.append("")
        lines.append("1. For Google Fonts available above: automatically mapped ✓")
        lines.append("2. For custom/Canva fonts:")
        lines.append("   • Option A: Use the suggested web-safe alternative")
        lines.append("   • Option B: Find similar fonts on Google Fonts")
        lines.append("   • Option C: Upload custom fonts to Webflow")
        lines.append("")
        lines.append("3. To override mappings, edit font_mappings.json")
        lines.append("")

        return '\n'.join(lines)

    def _suggest_mapping(self, font_name: str) -> str:
        """Suggest a web font mapping"""
        # Import here to avoid circular dependency
        from font_mapper import FontMapper

        if font_name in FontMapper.FONT_MAPPINGS:
            mapping = FontMapper.FONT_MAPPINGS[font_name]
            if 'google' in mapping:
                return f"Google Font: {mapping['google']}"
            elif 'fallback' in mapping:
                return f"Web-safe fallback: {mapping['fallback']}"

        # Suggest alternatives for common patterns
        font_lower = font_name.lower()

        if 'sans' in font_lower or 'arial' in font_lower or 'helvetica' in font_lower:
            return "Suggested alternative: Roboto or Open Sans (Google Fonts)"
        elif 'serif' in font_lower or 'times' in font_lower or 'georgia' in font_lower:
            return "Suggested alternative: Merriweather or Playfair Display (Google Fonts)"
        elif 'mono' in font_lower or 'code' in font_lower or 'courier' in font_lower:
            return "Suggested alternative: Roboto Mono (Google Fonts)"
        elif 'script' in font_lower or 'handwriting' in font_lower:
            return "Suggested alternative: Dancing Script or Pacifico (Google Fonts)"
        elif 'display' in font_lower or 'decorative' in font_lower:
            return "Suggested alternative: Bebas Neue or Anton (Google Fonts)"
        else:
            return "⚠️  No automatic mapping - will use web-safe fallback"

    def export_fonts_list(self, output_path: str = "detected_fonts.json"):
        """Export list of detected fonts as JSON"""
        export_data = {
            'total_fonts': len(self.fonts_found),
            'fonts': {}
        }

        for font_name, data in self.fonts_found.items():
            export_data['fonts'][font_name] = {
                'usage_count': data['count'],
                'slides': sorted(data['slides']),
                'sizes': sorted([float(s) for s in data['sizes']]),
                'bold_count': data['bold_count'],
                'italic_count': data['italic_count'],
                'colors': sorted(data['colors'])
            }

        with open(output_path, 'w', encoding='utf-8') as f:
            json.dump(export_data, f, indent=2)

        print(f"✓ Font data exported to: {output_path}")

    def create_custom_mappings(self, output_path: str = "font_mappings_custom.json"):
        """Create a template for custom font mappings"""
        mappings = {}

        for font_name in self.fonts_found.keys():
            # Check if already mapped
            from font_mapper import FontMapper
            if font_name not in FontMapper.FONT_MAPPINGS:
                # Create placeholder mapping
                mappings[font_name] = {
                    "google": "",  # Fill in Google Font name if available
                    "fallback": "sans-serif",
                    "notes": "Custom mapping needed - replace with similar web font"
                }

        if mappings:
            with open(output_path, 'w', encoding='utf-8') as f:
                json.dump(mappings, f, indent=2)

            print(f"✓ Custom mappings template created: {output_path}")
            print(f"  Edit this file to add your font mappings")
        else:
            print("✓ All fonts already have mappings!")


def main():
    """Main CLI"""
    print("\n" + "=" * 70)
    print("  FONT INSPECTOR - Analyze PowerPoint Fonts")
    print("=" * 70 + "\n")

    if len(sys.argv) < 2:
        print("Usage: python font_inspector.py <powerpoint_file.pptx>")
        print("\nOptions:")
        print("  python font_inspector.py file.pptx              # Show report")
        print("  python font_inspector.py file.pptx --export     # Export JSON")
        print("  python font_inspector.py file.pptx --mappings   # Create mappings template")
        sys.exit(1)

    pptx_file = sys.argv[1]

    if not os.path.exists(pptx_file):
        print(f"❌ Error: File not found: {pptx_file}")
        sys.exit(1)

    # Create inspector
    inspector = FontInspector(pptx_file)

    # Run inspection
    report = inspector.inspect()

    if report:
        print(report)

        # Handle options
        if '--export' in sys.argv or '-e' in sys.argv:
            inspector.export_fonts_list()

        if '--mappings' in sys.argv or '-m' in sys.argv:
            inspector.create_custom_mappings()

        print("\n" + "=" * 70)
        print("Next steps:")
        print("  1. Review the fonts detected above")
        print("  2. For custom fonts, use --mappings to create a template")
        print("  3. Run the converter with: python canva_to_webflow.py file.pptx")
        print("=" * 70 + "\n")


if __name__ == "__main__":
    main()

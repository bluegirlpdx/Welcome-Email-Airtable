"""
PowerPoint Extractor Module
Extracts content, layout, and styling information from PowerPoint slides
"""

import os
import re
from typing import List, Dict, Any, Tuple
from io import BytesIO
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import Image


class SlideElement:
    """Represents an element extracted from a slide"""

    def __init__(self, element_type: str, content: Any, position: Dict[str, float],
                 styling: Dict[str, Any], z_index: int = 0):
        self.type = element_type  # 'text', 'image', 'shape'
        self.content = content
        self.position = position  # {'left', 'top', 'width', 'height'}
        self.styling = styling
        self.z_index = z_index

    def to_dict(self) -> Dict[str, Any]:
        """Convert to dictionary for easier processing"""
        return {
            'type': self.type,
            'content': self.content,
            'position': self.position,
            'styling': self.styling,
            'z_index': self.z_index
        }


class SlideData:
    """Represents data extracted from a single slide"""

    def __init__(self, slide_number: int, title: str = ""):
        self.slide_number = slide_number
        self.title = title
        self.elements: List[SlideElement] = []
        self.background_color = None
        self.width = 0
        self.height = 0

    def add_element(self, element: SlideElement):
        """Add an element to the slide"""
        self.elements.append(element)

    def to_dict(self) -> Dict[str, Any]:
        """Convert to dictionary"""
        return {
            'slide_number': self.slide_number,
            'title': self.title,
            'background_color': self.background_color,
            'width': self.width,
            'height': self.height,
            'elements': [el.to_dict() for el in self.elements]
        }


class PowerPointExtractor:
    """Extracts content from PowerPoint files"""

    def __init__(self, pptx_path: str, output_dir: str = "./output"):
        self.pptx_path = pptx_path
        self.output_dir = output_dir
        self.presentation = None
        self.slides_data: List[SlideData] = []
        self.images_dir = os.path.join(output_dir, "images")

        # Create output directories
        os.makedirs(output_dir, exist_ok=True)
        os.makedirs(self.images_dir, exist_ok=True)

    def load_presentation(self):
        """Load the PowerPoint presentation"""
        try:
            self.presentation = Presentation(self.pptx_path)
            print(f"✓ Loaded presentation: {self.pptx_path}")
            print(f"  Slides: {len(self.presentation.slides)}")
            return True
        except Exception as e:
            print(f"✗ Error loading presentation: {e}")
            return False

    def extract_all_slides(self) -> List[SlideData]:
        """Extract content from all slides"""
        if not self.presentation:
            if not self.load_presentation():
                return []

        for idx, slide in enumerate(self.presentation.slides, 1):
            print(f"\nProcessing slide {idx}...")
            slide_data = self.extract_slide(slide, idx)
            self.slides_data.append(slide_data)
            print(f"  ✓ Extracted {len(slide_data.elements)} elements")

        return self.slides_data

    def extract_slide(self, slide, slide_number: int) -> SlideData:
        """Extract content from a single slide"""
        slide_data = SlideData(slide_number)

        # Get slide dimensions
        slide_data.width = self.presentation.slide_width.inches
        slide_data.height = self.presentation.slide_height.inches

        # Extract background color if available
        slide_data.background_color = self._extract_background_color(slide)

        # Process shapes in order (for z-index)
        for z_idx, shape in enumerate(slide.shapes):
            element = self._extract_shape(shape, slide_number, z_idx)
            if element:
                slide_data.add_element(element)

                # Use first text box as title if no title set
                if not slide_data.title and element.type == 'text':
                    slide_data.title = self._clean_text(element.content)[:50]

        # Generate default title if still empty
        if not slide_data.title:
            slide_data.title = f"Section {slide_number}"

        return slide_data

    def _extract_shape(self, shape, slide_number: int, z_index: int) -> SlideElement:
        """Extract a shape from the slide"""
        try:
            # Get position
            position = {
                'left': shape.left.inches if hasattr(shape, 'left') else 0,
                'top': shape.top.inches if hasattr(shape, 'top') else 0,
                'width': shape.width.inches if hasattr(shape, 'width') else 0,
                'height': shape.height.inches if hasattr(shape, 'height') else 0
            }

            # Text box or shape with text
            if shape.has_text_frame:
                return self._extract_text_element(shape, position, z_index)

            # Image
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                return self._extract_image_element(shape, slide_number, position, z_index)

            # Other shapes
            elif shape.shape_type in [MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.FREEFORM]:
                return self._extract_shape_element(shape, position, z_index)

        except Exception as e:
            print(f"    Warning: Error extracting shape: {e}")

        return None

    def _extract_text_element(self, shape, position: Dict, z_index: int) -> SlideElement:
        """Extract text element with styling"""
        text_content = []
        styling = {
            'font_family': None,
            'font_size': None,
            'color': None,
            'bold': False,
            'italic': False,
            'alignment': 'left'
        }

        for paragraph in shape.text_frame.paragraphs:
            para_text = paragraph.text.strip()
            if para_text:
                # Get styling from first run if available
                if paragraph.runs:
                    run = paragraph.runs[0]
                    if run.font.name:
                        styling['font_family'] = run.font.name
                    if run.font.size:
                        styling['font_size'] = run.font.size.pt
                    if run.font.color and hasattr(run.font.color, 'rgb'):
                        styling['color'] = self._rgb_to_hex(run.font.color.rgb)
                    styling['bold'] = run.font.bold
                    styling['italic'] = run.font.italic

                # Get alignment
                if paragraph.alignment:
                    alignment_map = {1: 'left', 2: 'center', 3: 'right', 4: 'justify'}
                    styling['alignment'] = alignment_map.get(paragraph.alignment, 'left')

                text_content.append(para_text)

        content = '\n'.join(text_content)
        return SlideElement('text', content, position, styling, z_index)

    def _extract_image_element(self, shape, slide_number: int,
                              position: Dict, z_index: int) -> SlideElement:
        """Extract image and save to file"""
        try:
            image = shape.image
            image_bytes = image.blob

            # Determine file extension
            ext = image.ext or 'png'
            filename = f"slide-{slide_number}-img-{z_index}.{ext}"
            filepath = os.path.join(self.images_dir, filename)

            # Save image
            with open(filepath, 'wb') as f:
                f.write(image_bytes)

            styling = {
                'alt': shape.name or f"Image {z_index}",
                'object_fit': 'cover'
            }

            return SlideElement('image', filename, position, styling, z_index)

        except Exception as e:
            print(f"    Warning: Error extracting image: {e}")
            return None

    def _extract_shape_element(self, shape, position: Dict, z_index: int) -> SlideElement:
        """Extract other shapes (rectangles, etc.)"""
        styling = {
            'shape_type': str(shape.shape_type),
            'fill_color': None,
            'border_color': None,
            'border_width': 0
        }

        # Try to extract fill color
        if shape.fill.type == 1:  # Solid fill
            try:
                if hasattr(shape.fill.fore_color, 'rgb'):
                    styling['fill_color'] = self._rgb_to_hex(shape.fill.fore_color.rgb)
            except:
                pass

        # Try to extract border
        if shape.line:
            try:
                styling['border_width'] = shape.line.width.pt if shape.line.width else 0
                if hasattr(shape.line.color, 'rgb'):
                    styling['border_color'] = self._rgb_to_hex(shape.line.color.rgb)
            except:
                pass

        return SlideElement('shape', 'rectangle', position, styling, z_index)

    def _extract_background_color(self, slide) -> str:
        """Extract slide background color"""
        try:
            if slide.background.fill.type == 1:  # Solid fill
                if hasattr(slide.background.fill.fore_color, 'rgb'):
                    return self._rgb_to_hex(slide.background.fill.fore_color.rgb)
        except:
            pass
        return None

    @staticmethod
    def _rgb_to_hex(rgb) -> str:
        """Convert RGB tuple to hex color"""
        try:
            if rgb:
                return f"#{rgb[0]:02x}{rgb[1]:02x}{rgb[2]:02x}"
        except:
            pass
        return None

    @staticmethod
    def _clean_text(text: str) -> str:
        """Clean and normalize text"""
        # Remove extra whitespace
        text = re.sub(r'\s+', ' ', text)
        return text.strip()

    def get_slides_data(self) -> List[Dict[str, Any]]:
        """Get all slides data as dictionaries"""
        return [slide.to_dict() for slide in self.slides_data]


if __name__ == "__main__":
    # Example usage
    import sys

    if len(sys.argv) < 2:
        print("Usage: python pptx_extractor.py <powerpoint_file>")
        sys.exit(1)

    pptx_file = sys.argv[1]
    extractor = PowerPointExtractor(pptx_file)

    slides = extractor.extract_all_slides()

    print(f"\n{'='*50}")
    print(f"Extraction Complete!")
    print(f"{'='*50}")
    print(f"Total slides processed: {len(slides)}")
    print(f"Output directory: {extractor.output_dir}")

    # Print summary
    for slide in slides:
        print(f"\nSlide {slide.slide_number}: {slide.title}")
        print(f"  Elements: {len(slide.elements)}")
        for el in slide.elements:
            print(f"    - {el.type}: {str(el.content)[:50]}")
